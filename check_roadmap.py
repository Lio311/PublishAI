from docx import Document
doc = Document('PublishAI_Product_Roadmap.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
if len(doc.tables) > 0:
    for row in doc.tables[0].rows:
        print([cell.text for cell in row.cells])
print("Last 10 paragraphs:")
for p in doc.paragraphs[-10:]:
    print(p.text)
