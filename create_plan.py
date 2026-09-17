from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

document = Document()
document.core_properties.title = "תוכנית עבודה: פרויקט PublishAI"

# Set default font properties for RTL
style = document.styles['Normal']
style.font.name = 'Arial'
style.font.rtl = True

# Add a title
title = document.add_heading("תוכנית עבודה: פרויקט PublishAI - אוטומציית כתיבה והגשת מאמרים", 0)
title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Intro
p = document.add_paragraph("תוכנית עבודה זו מגדירה את השלבים, המשאבים, לו\"ז והוולידציות הנדרשים להשלמת פרויקט PublishAI. התוכנית התעדכנה לאחר מחקר Unknown Unknowns וכוללת התמודדות עם מגבלות מדיניות, חומות תשלום, ניתוח נתונים, אסטרטגיית דחיות וסליקה.")
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Team Section
h_team = document.add_heading("צוות הפרויקט", level=1)
h_team.alignment = WD_ALIGN_PARAGRAPH.RIGHT

p_team = document.add_paragraph()
p_team.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_team.add_run("מור וליאור: ").bold = True
p_team.add_run("מייסדי הפרויקט (Founders). מור מובילה עסקית וכלכלית ויועצת בנושא חוקי כתיבת מאמרים אקדמיים ומכתבים.\n")
p_team.add_run("דניאל: ").bold = True
p_team.add_run("מהנדסת תוכנה ואוטומציה (Software & Automation Engineer). אחראית על כתיבת הקוד ופיתוח המודולים במערכת.\n")
p_team.add_run("יעלי: ").bold = True
p_team.add_run("יועצת מקצועית (Advisory) לפיתוח ולדאטה.")

def add_stage(doc, title_text, goal, schedule, actions, dev_reqs, validation):
    h = doc.add_heading(title_text, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("מטרה: ").bold = True
    p.add_run(goal)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("לו\"ז מוערך: ").bold = True
    p.add_run(schedule)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("פעולות מרכזיות:").bold = True
    for action in actions:
        pa = doc.add_paragraph(action, style='List Bullet')
        pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("דרישות פיתוחיות וטכנולוגיות:").bold = True
    for req in dev_reqs:
        pa = doc.add_paragraph(req, style='List Bullet')
        pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("ולידציה:").bold = True
    for val in validation:
        pa = doc.add_paragraph(val, style='List Bullet')
        pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Stage 1
add_stage(document,
          "שלב 1: מחקר, מדיניות, סטטסיטיקה וכתיבת המאמר",
          "בדיקת מדיניות חוקית של הז'ורנלים כלפי AI, אפיון ה-20 ז'ורנלים, מחקר ספרות מלא, ניתוח אוטומטי, וכתיבה שאינה מזוהה כ-AI.",
          "4-6 שבועות.",
          [
              "מחקר מדיניות (Policy Check): וידוא האם העיתונים אוסרים על כתיבה/סיוע באמצעות AI, ומה דרישות הגילוי הנאות (Disclosure) שלהם.",
              "אפיון מקצה לקצה (בהנחיית מור): מיפוי 20 עיתונים, כללי ציטוט, דרישות מבוא ו-Cover letter.",
              "איסוף נתונים: גישה למאמרים מלאים (Full-text) תוך עקיפת Paywalls (באמצעות Proxy באפיון יעלי).",
              "ניתוח נתונים (Stats): ה-AI קורא קבצי CSV ומריץ עליהם קוד סטטיסטי בסביבת Sandbox לחילוץ מסקנות.",
              "כתיבת המאמר ויצירת גרפים באמצעות פייתון."
          ],
          [
              "דניאל תפתח סקריפטים מבוססי Pandoc/python-docx שימירו את הפלט ישירות לפורמט ה-Word הספציפי.",
              "הקמת סביבת Python Sandbox מאובטחת שבה ה-AI יכול להריץ קוד סטטיסטי.",
              "שילוב פסקת 'גילוי נאות' גמישה בפרומפטים במידה והז'ורנל דורש הצהרה על שימוש ב-AI."
          ],
          ["בדיקת אנטי-פלגיאט (Turnitin).", "וידוא משפטי/תקנוני שאנחנו לא מפרים את מדיניות ה-AI של העיתון המוגש."]
          )

# Stage 2
add_stage(document,
          "שלב 2: התממשקות להגשה אוטומטית (Automated Submission)",
          "חיבור אוטומטי (RPA) למערכת ההגשה, בצורה המחקה התנהגות אנושית, עם יכולת לעצירה והתערבות אנושית.",
          "3-4 שבועות.",
          ["בדיקת תנאי השימוש (ToS) של מערכות ההגשה (כגון Editorial Manager) בנוגע לשימוש בבוטים, והתאמת קצב הריצה כדי לא לקבל חסימה (Ban).",
           "פיתוח מנגנון הזדהות, מילוי טפסים, והעלאת מסמכים.",
           "מנגנון 'Pause & Resume': הבוט נעצר במסכים מסובכים, שולח התראה, ממתין שהמשתמש יזין את המידע, וממשיך.",
           "עצירה לתשלום (APC): הבוט שולח לינק והמשתמש מבצע את התשלום ידנית."],
          ["פיתוח RPA על ידי דניאל (Selenium/Playwright) הכולל השהיות אקראיות (Human-like delays) כדי לא לעורר חשד של מנגנוני Anti-Bot.", 
           "מערכת התראות משתמש (Slack/Email/Web UI) להתערבות אנושית."],
          ["ביצוע 20 הגשות ניסיון (טסטים) ללא קריסת המערכת או קבלת חסימת IP מהשרת."]
          )

# Stage 3
add_stage(document,
          "שלב 3: הפינג-פונג (Peer Review) ואסטרטגיית מפל (Cascade)",
          "ניהול ביקורות, תיקונים אוטומטיים, והגשות מחדש במקרה של דחייה.",
          "2-6 חודשים.",
          ["קליטה אוטומטית של החלטת העיתון (Decision letter).", 
           "ניתוח ביקורות סוקרים והכנת מכתב תגובה ותיקוני טקסט/גרפים בהתאמה.",
           "אסטרטגיית מפל (דחייה): במקרה של Reject, המערכת תציע למשתמש רשימת ז'ורנלים חלופיים. לאחר בחירת המשתמש, דניאל (המערכת) תפרמט הכל מחדש לפורמט החדש ותגיש."],
          ["מנגנון AI לניתוח ביקורות והצעת ז'ורנלים חלופיים (Recommendation Engine).",
           "סקריפטים להתאמה מחדש (Re-formatting) לז'ורנל חלופי בלחיצת כפתור."],
          ["וידוא אנושי למכתבי תגובה ולבחירת הז'ורנל החלופי לפני יריית ההגשה."]
          )

# Budget Section
h = document.add_heading("הערכת תקציב כלים ומודלי שפה (Claude API)", level=1)
h.alignment = WD_ALIGN_PARAGRAPH.RIGHT

p_budg = document.add_paragraph("התקציב לוקח בחשבון שריצות רבות ייעצרו או ייכשלו (Buffer פיתוח).")
p_budg.alignment = WD_ALIGN_PARAGRAPH.RIGHT

data = [
    ["סעיף / שלב", "תיאור פעילות", "הערכת קלט (טוקנים)", "הערכת פלט (טוקנים)"],
    ["שלב 1", "מחקר Full-text, מדיניות, סטטיסטיקה וקוד", "650,000", "50,000"],
    ["שלב 2", "אוטומציית הגשה (RPA) מחקה-אנוש", "150,000", "5,000"],
    ["שלב 3", "ניתוח ביקורות / המלצות Cascade", "200,000", "20,000"],
    ["סה\"כ לריצה", "מסלול ללא שגיאות", "1,000,000", "75,000"],
    ["סה\"כ (20 טסטים)", "כולל מקדם דיבוג (Buffer פי 2.5)", "50,000,000 (50M)", "3,750,000 (3.75M)"]
]

table = document.add_table(rows=1, cols=4)
table.style = 'Table Grid'
# Ensure table is RTL visual
table._tbl.tblPr.append(OxmlElement('w:bidiVisual'))

hdr_cells = table.rows[0].cells
for i in range(4):
    hdr_cells[i].text = data[0][i]

for row in data[1:]:
    row_cells = table.add_row().cells
    for i in range(4):
        row_cells[i].text = row[i]

p_total = document.add_paragraph("\nסיכום עלויות מוערך (חודש פיתוח):\n"
                                 "- Claude API (קלט/פלט מורחבים): כ-$200 בחודש.\n"
                                 "- מנוי Claude Code: $100 בחודש.\n"
                                 "- סה\"כ: כ-$300 בחודש לפיתוח.")
p_total.alignment = WD_ALIGN_PARAGRAPH.RIGHT


# -------------------------------------------------------------
# POST-PROCESSING: APPLY FULL RTL TO EVERY ELEMENT
# -------------------------------------------------------------
for paragraph in document.paragraphs:
    # 1. Align Right
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # 2. Add w:bidi to paragraph properties
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        pPr.append(OxmlElement('w:bidi'))
    # 3. Add rtl = True to every run
    for run in paragraph.runs:
        run.font.rtl = True
        run.font.cs_bold = run.font.bold

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                pPr = p._p.get_or_add_pPr()
                if pPr.find(qn('w:bidi')) is None:
                    pPr.append(OxmlElement('w:bidi'))
                for r in p.runs:
                    r.font.rtl = True
                    r.font.cs_bold = r.font.bold

document.save("PublishAI_Project_Plan.docx")
print("Updated document saved successfully.")
