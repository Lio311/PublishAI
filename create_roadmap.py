from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def make_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
    for run in paragraph.runs:
        run.font.rtl = True
        run.font.cs_bold = run.font.bold

document = Document()
document.core_properties.title = "מפרט מוצר ותוכנית עבודה - PublishAI"

style = document.styles['Normal']
style.font.name = 'Arial'
style.font.rtl = True

# Title
title = document.add_heading("מפרט מוצר ותוכנית עבודה - PublishAI", 0)
make_rtl(title)

# 1. Intro
h1 = document.add_heading("1. תקציר", level=1)
make_rtl(h1)
p1 = document.add_paragraph("מערכת SaaS חכמה המקבלת מאמרים אקדמיים ונתונים סטטיסטיים, ומעבירה אותם תהליך אוטומטי-למחצה (היברידי) להפיכתם למאמרים מוכנים לפרסום בכתבי עת מובילים. המערכת מתבססת על שרשרת של סוכני AI מומחים, סביבות הרצת קוד מבודדות, אוטומציות להגשה \"מחקה אנוש\", וטיפול בדחיות או ביקורות עמיתים.")
make_rtl(p1)

# 2. Team
h2 = document.add_heading("2. צוות הפרויקט", level=1)
make_rtl(h2)
p2 = document.add_paragraph()
make_rtl(p2)
p2.add_run("מור וליאור: ").bold = True
p2.add_run("מייסדי הפרויקט (Founders). מור מובילה עסקית ויועצת בחוקי כתיבת מאמרים אקדמיים. ליאור מוביל את הטכנולוגיה מאחורי המוצר.\n")
p2.add_run("דניאל: ").bold = True
p2.add_run("מתכנתת ומבצעת משימות טכנולוגיות של הפיתוח.\n")
p2.add_run("יעלי: ").bold = True
p2.add_run("יועצת מקצועית לפיתוח ולדאטה.")

# 3. Features
h_feat = document.add_heading("3. חוויית המשתמש והפיצ'רים המרכזיים במוצר", level=1)
make_rtl(h_feat)

features = [
    ("דאשבורד העלאה וראיון מקדים (Upload & Clarification): ", "מסך גרירה לקבצי מקור ושאלון אמריקאי קצר שבאמצעותו ה-AI מבין את דרישות המשתמש."),
    ("פאנל עריכה (DiffEditor / Track Changes): ", "ממשק המציג את השינויים של ה-AI ומאפשר אישור או דחייה ברמת הסקציה (בדומה ל-Word)."),
    ("דאשבורד תוצרים (Artifact Dashboard): ", "ריכוז התוצרים הסופיים - המאמר המעוצב, מכתב ה-Cover Letter, ודוח ביקורת עמיתים סימולטיבי."),
    ("סביבת נתונים (Data Science Sandbox): ", "פיצ'ר מתקדם בו המערכת קוראת קבצי CSV, מריצה קוד סטטיסטי (Python) ומשלבת P-values וגרפים במאמר."),
    ("מודל הערות עמיתים (Reviewer Comments Modal): ", "מסך להזנת ביקורות עורכים (ממייל/PDF) ליצירת מכתב תגובה (Rebuttal) וגרסה מתוקנת.")
]
for f_title, f_desc in features:
    p = document.add_paragraph(style='List Bullet')
    make_rtl(p)
    r1 = p.add_run(f_title)
    r1.bold = True
    p.add_run(f_desc)

# 4. Tech Stack
h3 = document.add_heading("4. ארכיטקטורה טכנולוגית (Tech Stack)", level=1)
make_rtl(h3)
tech_stack = [
    "Frontend & Backend: Next.js 14/16 (App Router), React 19, TypeScript, Tailwind CSS.",
    "Database & Storage: Neon Postgres (Serverless DB), Drizzle ORM, Vercel Blob.",
    "Orchestration: Inngest (ניהול סוכנים ברקע).",
    "AI Engine: Vercel AI SDK, Claude 3.5 Sonnet/Opus, LangGraph.",
    "Data Science: E2B Code Interpreter (להרצת קוד פייתון מבודד)."
]
for item in tech_stack:
    p = document.add_paragraph(item, style='List Bullet')
    make_rtl(p)

# 5. Phased Implementation
h4 = document.add_heading("5. תוכנית עבודה ופיתוח (מחולקת ל-3 שלבים מרכזיים)", level=1)
make_rtl(h4)
p_intro = document.add_paragraph("התוכנית מאחדת את הפיתוח אל תוך 3 שלבי המוצר המרכזיים - כתיבה, הגשה, וביקורת עמיתים.")
make_rtl(p_intro)

def add_stage(title_text, items):
    h = document.add_heading(title_text, level=2)
    make_rtl(h)
    for item in items:
        p = document.add_paragraph(item, style='List Bullet')
        make_rtl(p)

add_stage("שלב 1: מחקר, סטטיסטיקה וכתיבת המאמר", [
    "תשתיות ליבה: הקמת פרויקט Next.js, מסד נתונים (Neon), אימות מאובטח (NextAuth) וניהול העלאות.",
    "שרשרת 9 סוכני AI (Inngest): יצירת מערך סוכנים מ-Clarification עד לכתיבה אקדמית (Opus/Sonnet).",
    "חיבור לספרות ועקיפת Paywalls: אינטגרציה עם PubMed/arXiv ושימוש ב-Proxy אקדמי/Scraping להבאת מאמרים מלאים (Full-text).",
    "סטטיסטיקה מבודדת (E2B): ה-AI קורא CSV, מריץ Python לחילוץ P-values ומסקנות, ומשלב גרפים במאמר.",
    "רגולציה ו-Policy Check: מערכת בדיקת AI-Policy מול העיתון, ושתילת פסקת \"גילוי נאות\" במידת הצורך.",
    "עיצוב מחמיר: שימוש ב-Pandoc/python-docx ליציקת המאמר והתוצרים לתבנית Word התואמת 1:1 לדרישות כתב העת."
])

add_stage("שלב 2: התחברות לז'ורנל והגשה (Submission)", [
    "חיבור API לז'ורנלים: אינטגרציה ישירה ו-Native למערכות תומכות (כגון OJS ו-WordPress).",
    "הגשה אוטומטית (Anti-Ban RPA): פיתוח אוטומציית דפדפן (Selenium/Playwright) להזנת טפסים לעיתונים מורכבים. כולל זמני השהייה אקראיים כדי לדמות התנהגות אנושית ולמנוע חסימה.",
    "עצירה לתשלום (Pause for APC): מנגנון העוצר את תהליך ההגשה אם נדרש תשלום דמי פרסום, משגר התראה למשתמש, וממתין עד לסליקה ידנית לשם המשך."
])

add_stage("שלב 3: פינג פונג הגשות (Peer Review & Cascade)", [
    "סוכן ביקורות (RebuttalAgent): ניתוח אוטומטי של הערות Reviewers שהתקבלו במייל, הצעת תגובות, עדכון המאמר, והפקת מכתב תגובה (Point-by-point).",
    "הגשה מחדש (Resubmission): עדכון הגרסאות ב-Database ושליחת חבילת ההגשה המתוקנת (v2) לז'ורנל.",
    "אסטרטגיית מפל בדחייה (Cascade): במקרה של דחייה סופית (Reject), המערכת תמליץ על 3 ז'ורנלים חלופיים, תבצע Re-formatting אוטומטי לעיצוב החדש, ותשגר הגשה חדשה ללא מאמץ מהמשתמש."
])

document.save("PublishAI_Product_Roadmap.docx")
print("Product Roadmap document saved successfully.")
