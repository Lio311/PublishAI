from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def make_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        pPr.append(OxmlElement('w:bidi'))
    for run in paragraph.runs:
        run.font.rtl = True
        run.font.cs_bold = run.font.bold

# Load the user's existing styled document
doc_path = 'PublishAI_Project_Plan.docx'
try:
    document = Document(doc_path)
except Exception as e:
    print(f"Error opening document: {e}")
    exit(1)

# Append a separator or title for the new phases
p_sep = document.add_paragraph("\n=========================================\n")
p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER

h_new = document.add_heading("השלבים החדשים (Unknown Unknowns - שלבים 6-8)", level=1)
h_new.alignment = WD_ALIGN_PARAGRAPH.RIGHT
make_rtl(h_new)

def add_stage(doc, title_text, description, bullets):
    h = doc.add_heading(title_text, level=2)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    make_rtl(h)
    
    p = doc.add_paragraph(description)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    make_rtl(p)
    
    for bullet in bullets:
        pb = doc.add_paragraph(bullet, style='List Bullet')
        pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        make_rtl(pb)

# Phase 6
add_stage(document,
          "שלב 6: מחקר מתקדם, דאטה ומדיניות AI",
          "השלמת פערים מול חומות תשלום, ניתוח נתונים, ובדיקות תקנון אקדמי.",
          [
              "מדיניות AI (Policy Check): בדיקה אוטומטית האם העיתון דורש גילוי נאות (Disclosure) על שימוש ב-AI ושתילת הפסקה במאמר במידת הצורך.",
              "עקיפת חומות תשלום (Paywalls): חיבור ל-Proxy אקדמי או פיתוח Advanced Scraping כדי לחלץ מאמרים מלאים (Full-text) ולא רק אבסטרקטים.",
              "סטטיסטיקה מורחבת בסנדבוקס: אינטגרציה עמוקה עם E2B Code Interpreter המאפשרת ל-AI לקרוא קבצי CSV, להריץ Python לחישוב P-values ולייצר גרפים."
          ])

# Phase 7
add_stage(document,
          "שלב 7: אוטומציית הגשות 'מחקה אנוש' (RPA Anti-Bot)",
          "בניית בוטים להגשה אוטומטית שלא נחסמים על ידי מערכות אבטחה של כתבי עת.",
          [
              "בוט מושהה (Human-like RPA): פיתוח באמצעות Selenium/Playwright עם זמני השהייה אקראיים כדי לעקוף חסימות (Ban) במערכות כגון Editorial Manager.",
              "עצירה לתשלום (Pause for APC): עצירת האוטומציה ושליחת התראה למשתמש לבצע סליקה ידנית במידה ונדרש תשלום."
          ])

# Phase 8
add_stage(document,
          "שלב 8: אסטרטגיית מפל ועיצוב מחמיר (Cascade & Strict Formatting)",
          "טיפול חכם בדחיות (Reject) ואריזה גרפית מחמירה לפי דרישות כתב העת.",
          [
              "אסטרטגיית מפל (Cascade): במקרה של דחייה סופית, המערכת תמליץ על 3 ז'ורנלים חלופיים ותבצע Re-formatting למאמר לפי חוקי הז'ורנל החדש שייבחר.",
              "עיצוב מחמיר למקור: שילוב סקריפטים (Pandoc / python-docx) שמייצאים את התוצר הסופי בדיוק לתבנית ה-Word הספציפית של כתב העת."
          ])

try:
    document.save(doc_path)
    print("Successfully appended to the existing document without breaking original styling.")
except Exception as e:
    print(f"Error saving document: {e}")
