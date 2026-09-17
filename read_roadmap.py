from docx import Document
doc = Document('PublishAI_Product_Roadmap.docx')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text}")
for i, t in enumerate(doc.tables):
    print(f"T{i}: Table with {len(t.rows)} rows")
