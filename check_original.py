from docx import Document
doc = Document('PublishAI_Project_Plan.docx')
print(f"Total paragraphs in Project_Plan: {len(doc.paragraphs)}")
print("Last 10 paragraphs:")
for p in doc.paragraphs[-10:]:
    print(p.text)
